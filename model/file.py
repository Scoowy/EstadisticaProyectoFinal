#!/usr/bin/python
# -*- coding: utf-8 -*-
# Scoowy - Juan Gahona

import openpyxl
import pandas as pd

from openpyxl import Workbook
from openpyxl.chart import PieChart, Reference, BarChart
from openpyxl.utils.dataframe import dataframe_to_rows


class ExcelDoc(object):

    def __init__(self, name: str, df: pd.DataFrame):
        self.path = './data/'
        self.name = name
        self.df = df
        self.destFile = self.path + self.name

    def saveDoc(self):
        wb = Workbook()
        ws1 = wb.active
        ws1.title = 'Datos'

        for row in dataframe_to_rows(self.df, index=False, header=True):
            ws1.append(row)

        ws1['A10'] = 'TOTAL'
        ws1['B10'] = '=SUM(B2:B9)'
        ws1['D10'] = '=SUM(D2:D9)'
        ws1['E10'] = '=SUM(E2:E9)'

        for cell in ws1['E']:
            cell.style = 'Percent'

        for cell in ws1['A'] + ws1[1]:
            cell.style = 'Pandas'

        ws2 = wb.create_sheet(title='Grafico')

        bar = BarChart()
        bar.type = 'col'
        bar.style = 10
        bar.title = 'Frecuencia de colores usados'
        bar.y_axis.title = 'Frecuencia'
        bar.x_axis.title = 'Color'
        labels1 = Reference(ws1, min_col=1, min_row=2, max_row=9)
        data1 = Reference(ws1, min_col=2, min_row=2, max_row=9)
        bar.add_data(data1, titles_from_data=True)
        bar.set_categories(labels1)
        bar.shape = 4
        ws2.add_chart(bar, 'A1')

        pie = PieChart()
        labels = Reference(ws1, min_col=1, min_row=2, max_row=9)
        data = Reference(ws1, min_col=5, min_row=2, max_row=9)
        pie.add_data(data, titles_from_data=True)
        pie.set_categories(labels)
        pie.title = 'Porcentaje de colores usados'

        ws2.add_chart(pie, 'A16')

        wb.save(filename='{}{}'.format(self.path, self.name))


class PandasExcel(object):

    def __init__(self, name: str, sheet: str):
        self.path = './data/'
        self.name = name
        self.sheet = sheet
        self.columnNames = ['Xi', 'fi']

    def openDoc(self):
        return pd.read_excel('{}{}'.format(
            self.path, self.name), self.sheet, names=self.columnNames)
